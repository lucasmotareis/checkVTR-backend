"""Gera o Word usando somente resultados recebidos; nunca inventa medicoes.

Requer python-docx. O PDF deve ser exportado desse mesmo Word e revisado.
"""
import argparse
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path

from docx import Document
from docx.image.image import Image
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def load_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig')) if path.exists() else {}


def values(data, name):
    return data.get('metrics', {}).get(name, {}).get('values', {})


def number(value, suffix='', digits=2):
    if value is None:
        return 'Pendente'
    return f'{value:,.{digits}f}'.replace(',', 'X').replace('.', ',').replace('X', '.') + suffix


def stage(data, level):
    count = values(data, f'measured_requests{{level:{level}}}').get('count')
    elapsed = values(data, f'stage_elapsed_seconds{{level:{level}}}').get('max')
    trend = values(data, f'client_duration_ms{{level:{level}}}')
    failure = values(data, f'unexpected_errors{{level:{level}}}').get('rate')
    return [number(trend.get('avg'), ' ms'), number(trend.get('p(95)'), ' ms'),
            number(count / elapsed if count is not None and elapsed else None, '/s'),
            number(failure * 100 if failure is not None else None, '%')]


def paragraph(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    for row_index, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        prevent_split = OxmlElement('w:cantSplit')
        props.append(prevent_split)
        if row_index == 0:
            props.append(OxmlElement('w:tblHeader'))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E8EDF2' if row_index == 0 else 'FFFFFF')
            cp.append(shading)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement('w:' + side)
                element.set(qn('w:w'), '90')
                element.set(qn('w:type'), 'dxa')
                margins.append(element)
            cp.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'bottom', 'left', 'right']:
                element = OxmlElement('w:' + edge)
                element.set(qn('w:val'), 'single')
                element.set(qn('w:sz'), '4')
                element.set(qn('w:color'), 'D9D9D9')
                borders.append(element)
            cp.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if i > 0 and len(headers) > 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_index == 0
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    paragraph(doc, '').paragraph_format.space_after = Pt(0)
    return t


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--results', type=Path, default=Path('observability/results'))
    parser.add_argument('--output', type=Path, default=Path('../../output/apm'))
    args = parser.parse_args()
    results = args.results.resolve()
    meta = load_json(results / 'experiment.json')
    baseline = load_json(results / 'baseline.json')
    apm = load_json(results / 'apm.json')
    unauthorized = load_json(results / 'unauthenticated.json')
    collection = meta.get('collection_date') or 'Ainda não realizada ou não informada'
    deployment = meta.get('deployment', {})
    apm_meta = meta.get('apm', {})
    dataset = meta.get('dataset', {})
    confirmed = meta.get('confirmed_homolog_url')
    for data, mode in [(baseline, 'baseline'), (apm, 'apm'), (unauthorized, 'unauthenticated')]:
        if data and (data.get('schema_version') != 1 or data.get('mode') != mode):
            raise ValueError(f'Resultado incompatível no modo {mode}')
        if data and (not confirmed or data.get('base_url') != confirmed):
            raise ValueError('O domínio dos resultados deve coincidir com confirmed_homolog_url em experiment.json.')
    complete = bool(meta.get('collection_date') and apm_meta.get('deployment_confirmed')
                    and apm_meta.get('privacy_verified') and deployment.get('database_isolated_confirmed')
                    and deployment.get('volume_isolated_confirmed') and dataset.get('synthetic_only')
                    and baseline.get('measurement_complete') and apm.get('measurement_complete')
                    and unauthorized.get('measurement_complete') and apm_meta.get('evidence_files')
                    and meta.get('findings'))
    doc = Document()
    # Remover bordas herdadas do template, inclusive a linha azul do estilo Title.
    for element in list(doc.styles.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2)
    section.header_distance = section.footer_distance = Cm(0.8)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Title'].paragraph_format.space_after = Pt(9)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(11)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    doc.styles['Heading 2'].font.size = Pt(12)
    header = section.header.paragraphs[0]
    header.text = 'CheckVTR | Monitoramento em homologação'
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Página ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = 'Monitoramento APM do CheckVTR em homologação'
    doc.core_properties.author = 'Equipe CheckVTR'
    doc.core_properties.subject = 'Atividade de monitoramento de aplicações com New Relic'

    paragraph(doc, 'Monitoramento APM do CheckVTR em homologação', 'Title')
    paragraph(doc, 'Relatório da atividade com New Relic', 'Subtitle')
    status = 'Coleta registrada com evidências' if complete else 'Versão preliminar com coleta de homologação pendente'
    p = paragraph(doc, status)
    p.runs[0].bold = True
    paragraph(doc, 'Preparação e validações locais realizadas em 7 de outubro de 2026. '
              f'Data da coleta no servidor: {collection}. Fuso de análise: America/Sao_Paulo.')
    participants = meta.get('participants') or []
    paragraph(doc, 'Participantes: ' + (', '.join(participants) if participants else 'ainda não informados.')
              + (f' Disciplina: {meta["course"]}.' if meta.get('course') else ''))
    doc.add_heading('Objetivo e situação da atividade', level=1)
    paragraph(doc, 'O CheckVTR é um sistema de inspeção de viaturas da Polícia Militar. '
              'Esta atividade prepara o monitoramento do backend implantado em homologação '
              'para observar a duração das requisições, a frequência de erros, o uso da JVM '
              'e o tempo das operações de acesso ao banco de dados.')
    paragraph(doc, 'A integração e os testes locais foram concluídos. A conta New Relic, '
              'a configuração dos segredos e o deploy no servidor são operados pelo responsável '
              'pela homologação. Enquanto essas etapas e a coleta não forem documentadas, '
              'não há evidência de desempenho do servidor nem comprovação de ingestão no painel.'
              if not complete else 'Os resultados apresentados foram extraídos dos arquivos de coleta '
              'e das capturas fornecidas pelo responsável pela homologação. As verificações '
              'locais são identificadas separadamente das medições do servidor.')
    doc.add_heading('Aplicação e arquitetura', level=1)
    paragraph(doc, 'O backend utiliza Spring Boot 3.5.5, Java 21, PostgreSQL e autenticação JWT. '
              'A organização segue Controller → Service → Repository, com acesso aos dados '
              'restrito ao batalhão do usuário. A aplicação é implantada em Docker, por Coolify, '
              'com Traefik no encaminhamento HTTPS. O usuário informou que homologação funciona '
              'em servidor separado de produção.')
    paragraph(doc, 'Fluxo observado: cliente → Traefik → backend Spring Boot → PostgreSQL. '
              'O agente New Relic opera dentro da JVM do backend e envia telemetria ao serviço '
              'APM por HTTPS. O frontend e a infraestrutura do host não são instrumentados nesta etapa.')

    doc.add_page_break()
    doc.add_heading('Integração da ferramenta APM', level=1)
    paragraph(doc, 'Foi incorporado à imagem o agente Java New Relic 9.4.0, obtido da fonte oficial '
              'com verificação SHA256 do arquivo baixado. A versão fixa permite reproduzir a '
              'configuração utilizada. A aplicação será identificada no painel como checkvtr-api-homolog.')
    table(doc, ['Configuração', 'Comportamento'], [
        ['APM_ENABLED=false', 'Padrão e reversão; JVM inicia sem carregar o agente.'],
        ['APM_ENABLED=true', 'Exige perfil homolog, ausência de perfil prod/production e chave de ingestão.'],
        ['NEW_RELIC_LICENSE_KEY', 'Segredo de execução no Coolify; não gravado no código ou na imagem.'],
        ['checkvtr-api-homolog', 'Identificação exclusiva da aplicação no APM.'],
    ], [5.2, 11.8])
    doc.add_heading('Privacidade e isolamento', level=1)
    paragraph(doc, 'A configuração desabilita o encaminhamento de logs da aplicação e o monitoramento '
              'do navegador. Exclui atributos de parâmetros, headers, cookies e identificação de '
              'usuários, remove mensagens de exceções e mantém consultas SQL ofuscadas. '
              'O experimento utiliza somente contas e dados fictícios. A aplicação, suas respostas '
              'e a autenticação JWT preservam os contratos existentes.')
    paragraph(doc, 'O Compose mantém o banco e os mounts existentes de homologação. Antes de publicar, '
              'é necessário confirmar no Coolify o recurso correto, a branch homolog, o perfil '
              'ativo e o isolamento do PostgreSQL e do volume. Os guias locais apresentam dois '
              'candidatos de domínio; nenhum deles deve ser tratado como confirmado apenas '
              'pela documentação. Domínio confirmado: ' + (confirmed or 'pendente de verificação.'))
    doc.add_heading('Execução e reversão', level=1)
    paragraph(doc, 'Primeiro, implantar a integração com APM desabilitado e coletar a referência. '
              'Depois, fornecer a chave na execução, habilitar a flag somente na homologação '
              'e reimplantar o serviço app. Confirmar no painel a aplicação e as transações '
              'antes de executar a carga com APM. A saída HTTPS do agente deve estar disponível.')
    paragraph(doc, 'Se ocorrer falha de inicialização ou regressão funcional, alterar APM_ENABLED '
              'para false e reimplantar somente app. Validar novamente login e listagem. '
              'A reversão não exige alteração ou remoção de banco, credenciais ou volumes.')

    doc.add_page_break()
    doc.add_heading('Metodologia do experimento', level=1)
    paragraph(doc, 'O k6 é executado no computador do operador. O script aceita somente os dois '
              'domínios candidatos de homologação e exige confirmação do destino. Os redirecionamentos '
              'HTTP são recusados. Matrícula e senha fictícias entram por variáveis temporárias; '
              'o token fica em memória e não é exportado nos resultados.')
    paragraph(doc, 'O login acontece antes da carga. Cada usuário executa consultas a viaturas, '
              'checklists paginados, dashboard e feed, seguidas por pausa de um segundo. '
              'As consultas não criam registros. O mesmo conjunto fictício deve ser mantido '
              'durante a referência e a coleta com agente.')
    table(doc, ['Etapa', 'Agente', 'Usuários', 'Duração'], [
        ['Aquecimento de cada coleta', 'Conforme coleta', '1', '2 minutos'],
        ['Referência', 'Desligado', '1', '5 minutos'],
        ['Medição com APM', 'Ligado', '1', '5 minutos'],
        ['Medição com APM', 'Ligado', '5', '5 minutos'],
        ['Medição com APM', 'Ligado', '10', '5 minutos'],
        ['Acesso sem autenticação', 'Ligado', '1 chamada', 'Execução separada'],
    ], [7, 3.5, 2.5, 4])
    paragraph(doc, 'Após um minuto de cada etapa, o teste avalia periodicamente a interrupção '
              'se os erros inesperados ultrapassarem 5% ou o p95 ultrapassar cinco segundos. '
              'Durante a carga autenticada, respostas 401 ou 403 encerram imediatamente a execução. '
              'Entre as etapas existem intervalos de dez segundos para finalizar chamadas anteriores.')
    doc.add_heading('Métricas e condições de comparação', level=1)
    paragraph(doc, 'A média e o p95 do cliente excluem login e aquecimento. O RPS por etapa é '
              'estimado pela quantidade de requisições dividida pelo tempo até a última sequência '
              'concluída. Também são registradas respostas HTTP e erros inesperados. No New Relic '
              'serão observados duração das transações, throughput, erros, heap, garbage collection '
              'e operações JDBC. A latência do cliente inclui rede e proxy e não deve ser '
              'confundida com a duração interna da transação.')
    server = meta.get('server', {})
    paragraph(doc, f'Condições registradas: CPU {server.get("cpu") or "pendente"}; '
              f'memória {server.get("memory_gib") or "pendente"} GiB; '
              f'sistema operacional {server.get("operating_system") or "pendente"}; '
              f'commit {deployment.get("commit") or "pendente"}. '
              f'Dados fictícios: {dataset.get("users") if dataset.get("users") is not None else "pendente"} usuários, '
              f'{dataset.get("vehicles") if dataset.get("vehicles") is not None else "pendente"} viaturas e '
              f'{dataset.get("checklists") if dataset.get("checklists") is not None else "pendente"} checklists.')

    doc.add_page_break()
    doc.add_heading('Resultados e evidências', level=1)
    paragraph(doc, 'As tabelas distinguem verificações locais de desempenho em homologação. '
              'Campos pendentes não representam zero erros ou tempo zero; indicam ausência '
              'de uma medição recebida. Os testes simulados não foram usados como resultados do servidor.')
    table(doc, ['Coleta', 'Média', 'p95', 'RPS', 'Erros'], [
        ['Referência 1 usuário'] + stage(baseline, '1'),
        ['APM 1 usuário'] + stage(apm, '1'),
        ['APM 5 usuários'] + stage(apm, '5'),
        ['APM 10 usuários'] + stage(apm, '10'),
    ], [5.4, 3, 3, 2.8, 2.8])
    for data, label in [(baseline, 'Referência'), (apm, 'Com APM')]:
        if data:
            paragraph(doc, f'{label}: início UTC {data.get("started_at_utc", "não informado")}; '
                      f'fim UTC {data.get("finished_at_utc", "não informado")}. '
                      f'Coleta completa: {"sim" if data.get("measurement_complete") else "não"}. '
                      f'Critérios atendidos: {"sim" if data.get("thresholds_passed") else "não"}.')
    blocked_status = ', '.join(code for code in ['401', '403'] if values(unauthorized, f'responses_{code}').get('count', 0))
    paragraph(doc, 'Chamada sem token: ' + (f'HTTP {blocked_status}, registrado separadamente.' if blocked_status
              else 'pendente de execução em homologação. No teste local simulado, o bloqueio foi verificado.'))
    table(doc, ['Verificação local', 'Resultado de 07/10/2026'], [
        ['Testes existentes do backend', '124 passaram com Java 21'],
        ['Empacotamento do JAR', 'Concluído'],
        ['Configuração Docker Compose', 'Validada com segredos fictícios'],
        ['Inicialização com e sem agente', 'Concluída com H2 em memória; HTTP 401 em /auth/me'],
        ['Proteções de inicialização', '7 cenários passaram'],
        ['Proteções da configuração k6', '5 testes passaram'],
        ['k6 contra HTTPS simulado local', '5 cenários passaram sem exportar credenciais'],
        ['Build da imagem Docker', 'Pendente; daemon Docker local indisponível'],
    ], [8.3, 8.7])
    paragraph(doc, 'A inicialização local utilizou chave fictícia e coletor direcionado para '
              '127.0.0.1:9. Ela verifica compatibilidade de inicialização, sem comprovar ingestão '
              'de telemetria na conta New Relic. O build e o deploy de homologação devem ser verificados no Coolify.')

    files = apm_meta.get('evidence_files', [])
    if files:
        for index, relative in enumerate(files, 1):
            path = (results / relative).resolve()
            if not path.is_relative_to(results) or not path.is_file():
                raise ValueError('As capturas precisam existir dentro da pasta de resultados.')
            doc.add_page_break()
            doc.add_heading(f'Evidência {index}', level=1)
            paragraph(doc, 'Captura fornecida da coleta em homologação. Conferir o período, '
                      'a identificação da aplicação e a ausência de dados sensíveis.')
            image = Image.from_file(str(path))
            width_cm = min(17, 20 * image.px_width / image.px_height)
            doc.add_picture(str(path), width=Cm(width_cm))
            paragraph(doc, f'Arquivo de evidência: {relative}.', 'Caption')

    doc.add_page_break()
    doc.add_heading('Análise e limitações', level=1)
    findings = meta.get('findings') or []
    if findings:
        for finding in findings:
            paragraph(doc, str(finding))
    else:
        paragraph(doc, 'Ainda não há medições ou capturas de homologação recebidas para identificar '
                  'a transação mais lenta, um gargalo no banco ou a evolução da memória. '
                  'Essas conclusões serão escritas somente após a execução e a análise dos dados reais. '
                  'A ausência de erros será um resultado válido se demonstrada na janela de coleta.')
    if apm_meta.get('observations'):
        paragraph(doc, apm_meta['observations'])
    if not files:
        paragraph(doc, 'Capturas pendentes: resumo APM, transações, operações de banco, JVM e erros. '
                  'Deverão ser inseridas no relatório com o mesmo período de medição e sem dados pessoais.')
    paragraph(doc, 'O experimento utiliza carga pequena, dados fictícios e uma janela curta. '
              'Cache, aquecimento da JVM e variação de rede podem afetar o resultado. '
              'Uma execução antes e depois não permite atribuir estatisticamente uma diferença '
              'ao custo do agente. O monitoramento JDBC não substitui análise interna do PostgreSQL. '
              'Os resultados não demonstram capacidade máxima e não devem ser extrapolados para produção.')
    doc.add_heading('Conclusão da etapa atual', level=1)
    paragraph(doc, 'A preparação disponibilizou instrumentação reversível, proteções para homologação '
              'e um procedimento reproduzível de coleta. As verificações locais confirmaram '
              'a inicialização do backend com e sem agente e a execução do teste contra respostas '
              'simuladas. A atividade de monitoramento no servidor permanece pendente até que '
              'o deploy, a ingestão no painel e os resultados reais sejam registrados.' if not complete
              else 'A integração permitiu registrar o comportamento do backend na janela de teste '
              'informada. A interpretação deve considerar os achados descritos nesta seção e '
              'os limites da carga utilizada. A configuração pode ser desligada por flag sem '
              'alterar o banco ou o volume de homologação.')
    doc.add_heading('Roteiro da apresentação', level=1)
    paragraph(doc, 'Sete minutos, sem slides: contexto e objetivo em um minuto; arquitetura e '
              'implantação em dois minutos; metodologia, painel e resultados em três minutos; '
              'conclusão e limitações em um minuto. As capturas inseridas no relatório permitem '
              'apresentar o experimento se o painel estiver indisponível.')
    doc.add_heading('Referências', level=1)
    references = [
        ('New Relic Agente Java 9.4.0', 'https://docs.newrelic.com/docs/release-notes/agent-release-notes/java-release-notes/java-agent-940/'),
        ('New Relic Instalação em Docker', 'https://docs.newrelic.com/docs/apm/agents/java-agent/additional-installation/install-new-relic-java-agent-docker/'),
        ('New Relic Configuração do agente', 'https://docs.newrelic.com/docs/apm/agents/java-agent/configuration/java-agent-configuration-config-file/'),
        ('Grafana k6 Resumo personalizado', 'https://grafana.com/docs/k6/latest/results-output/end-of-test/custom-summary/'),
        ('Grafana k6 Thresholds', 'https://grafana.com/docs/k6/latest/using-k6/thresholds/'),
    ]
    for title, url in references:
        p = paragraph(doc, title + '\n' + url)
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(9)
    args.output.mkdir(parents=True, exist_ok=True)
    output = args.output / 'Relatorio-APM-CheckVTR.docx'
    doc.save(output)
    print(output.resolve())
    print('Status: ' + status)


if __name__ == '__main__':
    main()
